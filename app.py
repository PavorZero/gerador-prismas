import streamlit as st
import pandas as pd
from pptx import Presentation
from pptx.util import Pt, Cm
import io
import os
import copy
from pptx.opc.constants import RELATIONSHIP_TYPE as RT


# ============================================================
# FUNÇÕES AUXILIARES
# ============================================================

def adjust_vertical_centering(slide, missing_fields):
    """Força a ordem NOME -> CARGO -> EMPRESA e centraliza no espaço original."""
    s_nome, s_cargo, s_empresa = None, None, None

    def find_shapes(shapes):
        nonlocal s_nome, s_cargo, s_empresa
        for s in shapes:
            if s.has_text_frame:
                txt = s.text_frame.text.upper()
                if "NOME" in txt: s_nome = s
                if "CARGO" in txt: s_cargo = s
                if "EMPRESA" in txt: s_empresa = s
            if s.shape_type == 6:  # GROUP
                find_shapes(s.shapes)

    find_shapes(slide.shapes)

    # Se não tiver as 3 caixas separadas, não há como forçar a ordem
    if not (s_nome and s_cargo and s_empresa):
        return
    if s_nome == s_cargo or s_nome == s_empresa or s_cargo == s_empresa:
        return

    # Calcula o centro vertical do bloco original
    y_min = min(s_nome.top, s_cargo.top, s_empresa.top)
    y_max = max(s_nome.top + s_nome.height, s_cargo.top + s_cargo.height, s_empresa.top + s_empresa.height)
    center_y = y_min + (y_max - y_min) / 2

    # Pega os espaçamentos originais entre as caixas para manter a proporção
    boxes = [s_nome, s_cargo, s_empresa]
    boxes.sort(key=lambda s: s.top) # Ordena temporariamente só para pegar o gap real
    gap1 = max(0, boxes[1].top - (boxes[0].top + boxes[0].height))
    gap2 = max(0, boxes[2].top - (boxes[1].top + boxes[1].height))
    avg_gap = (gap1 + gap2) / 2

    # Define a lista de caixas na ORDEM CORRETA (Nome, Cargo, Empresa)
    remaining = []
    if "NOME" not in missing_fields: remaining.append(s_nome)
    if "CARGO" not in missing_fields: remaining.append(s_cargo)
    if "EMPRESA" not in missing_fields: remaining.append(s_empresa)

    # Aplica o novo posicionamento mantendo o centro
    if len(remaining) == 3:
        total_h = s_nome.height + s_cargo.height + s_empresa.height + (2 * avg_gap)
        new_top = int(center_y - (total_h / 2))
        s_nome.top = new_top
        s_cargo.top = int(new_top + s_nome.height + avg_gap)
        s_empresa.top = int(new_top + s_nome.height + s_cargo.height + (2 * avg_gap))

    elif len(remaining) == 2:
        total_h = remaining[0].height + remaining[1].height + avg_gap
        new_top = int(center_y - (total_h / 2))
        remaining[0].top = new_top
        remaining[1].top = int(new_top + remaining[0].height + avg_gap)

    elif len(remaining) == 1:
        remaining[0].top = int(center_y - (remaining[0].height / 2))


def duplicate_slide(presentation, index):
    source = presentation.slides[index]
    dest = presentation.slides.add_slide(source.slide_layout)

    for shp in list(dest.shapes):
        shp._element.getparent().remove(shp._element)

    src_spTree = source.shapes._spTree
    dst_spTree = dest.shapes._spTree

    for child in list(dst_spTree):
        tag = child.tag.split('}')[-1]
        if tag not in ('nvGrpSpPr', 'grpSpPr'):
            dst_spTree.remove(child)

    for child in list(src_spTree):
        tag = child.tag.split('}')[-1]
        if tag in ('nvGrpSpPr', 'grpSpPr'):
            continue
        new_el = copy.deepcopy(child)
        dst_spTree.append(new_el)

    rels_map = {}
    for rel in source.part.rels.values():
        if rel.reltype == RT.SLIDE_LAYOUT:
            continue
        if 'notesSlide' in rel.reltype:
            continue
        if rel.is_external:
            new_rid = dest.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
        else:
            new_rid = dest.part.relate_to(rel.target_part, rel.reltype)
        rels_map[rel.rId] = new_rid

    for elem in dst_spTree.iter():
        for attr in list(elem.attrib.keys()):
            if attr.endswith('}embed') or attr.endswith('}link') or attr.endswith('}id'):
                old_rid = elem.get(attr)
                if old_rid and old_rid in rels_map:
                    elem.set(attr, rels_map[old_rid])

    return dest


def delete_slide(presentation, index):
    xml_slides = presentation.slides._sldIdLst
    slides = list(xml_slides)
    xml_slides.remove(slides[index])


def substitute_text_in_slide(slide, substituicoes, campos_para_limpar=None, apagar_shape=False, font_sizes=None):
    campos_para_limpar = campos_para_limpar or []
    font_sizes = font_sizes or {}
    
    # Força a ordem e centraliza
    adjust_vertical_centering(slide, campos_para_limpar)

    chaves_limpar_upper = [c.upper() for c in campos_para_limpar]

    def substituir_no_run(run, subs):
        texto = run.text
        for chave, valor in subs.items():
            if chave in texto or chave.capitalize() in texto or chave.lower() in texto:
                texto = texto.replace(chave, str(valor))
                texto = texto.replace(chave.capitalize(), str(valor))
                texto = texto.replace(chave.lower(), str(valor))
                
                if chave in font_sizes and font_sizes[chave]:
                    run.font.size = Pt(font_sizes[chave])
        run.text = texto

    def limpar_no_run(run, chaves):
        texto = run.text
        for chave in chaves:
            for variacao in (chave, chave.capitalize(), chave.lower()):
                if variacao in texto:
                    texto = texto.replace(variacao, "")
        run.text = texto

    def processar_text_frame(tf):
        for paragraph in tf.paragraphs:
            for run in paragraph.runs:
                substituir_no_run(run, substituicoes)
                limpar_no_run(run, chaves_limpar_upper)

    def processar_shape(shape):
        if shape.has_text_frame and chaves_limpar_upper:
            texto_shape = shape.text_frame.text.upper()
            contem_limpar = any(chave in texto_shape for chave in chaves_limpar_upper)

            if contem_limpar and apagar_shape:
                texto_restante = texto_shape
                for chave in chaves_limpar_upper:
                    texto_restante = texto_restante.replace(chave, "")
                if not texto_restante.strip(" \t\n\r-:|·•"):
                    shape._element.getparent().remove(shape._element)
                    return

        if shape.has_text_frame:
            processar_text_frame(shape.text_frame)

        if shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    processar_text_frame(cell.text_frame)

        if shape.shape_type == 6:  # GROUP
            for sub in list(shape.shapes):
                processar_shape(sub)

    for shape in list(slide.shapes):
        processar_shape(shape)


# ============================================================
# INTERFACE STREAMLIT
# ============================================================

st.set_page_config(page_title="Gerador de Prismas", page_icon="🔷", layout="centered")

st.title("🔷 Gerador de Prismas")
st.caption("Versão Web 1.7 - Criado por João Carlos")

st.markdown("---")

# PASSO 1: Upload do Excel
st.subheader("1️⃣ Arquivo Excel")
excel_file = st.file_uploader("Envie o arquivo Excel", type=["xlsx", "xls"])

if not excel_file:
    st.info("Aguardando o arquivo Excel para começar...")
    st.stop()

try:
    xls = pd.ExcelFile(excel_file)
    sheet_names = xls.sheet_names
    st.success(f"✅ Excel carregado: {len(sheet_names)} aba(s) encontrada(s)")
except Exception as e:
    st.error(f"Erro ao ler Excel: {e}")
    st.stop()

# PASSO 2: Escolher aba e Mapeamento
st.subheader("2️⃣ Selecionar Aba")
selected_sheet = st.selectbox("Aba do Excel", sheet_names)

df = pd.read_excel(excel_file, sheet_name=selected_sheet)
colunas = [''] + df.columns.tolist()

st.markdown("**Mapeamento de Colunas**")
col1, col2, col3 = st.columns(3)
with col1:
    nome_col = st.selectbox("Nome", colunas, index=0)
with col2:
    cargo_col = st.selectbox("Cargo (opcional)", colunas, index=0)
with col3:
    empresa_col = st.selectbox("Empresa", colunas, index=0)

# --- Configurações de Fonte Separadas ---
st.markdown("**Configurações de Texto e Layout**")
apagar_shape_vazio = st.checkbox("Apagar caixa de texto inteira se o campo não for usado", value=False, help="Se desmarcado, a caixa fica vazia, mas o texto que sobrar é centralizado automaticamente.")

font_sizes = {}
f_col1, f_col2, f_col3 = st.columns(3)

with f_col1:
    if nome_col:
        if st.checkbox("Alterar fonte do Nome"):
            font_sizes["NOME"] = st.number_input("Tamanho Nome (Pt)", min_value=8, max_value=150, value=28, step=1)

with f_col2:
    if cargo_col:
        if st.checkbox("Alterar fonte do Cargo"):
            font_sizes["CARGO"] = st.number_input("Tamanho Cargo (Pt)", min_value=8, max_value=150, value=18, step=1)

with f_col3:
    if empresa_col:
        if st.checkbox("Alterar fonte da Empresa"):
            font_sizes["EMPRESA"] = st.number_input("Tamanho Empresa (Pt)", min_value=8, max_value=150, value=18, step=1)

st.markdown("---")

# PASSO 3: Formato e Template
st.subheader("3️⃣ Formato e Template")

formato_prisma = st.radio(
    "Selecione o formato do prisma:",
    ["21x14,5 cm", "30x10,5 cm"],
    horizontal=True
)

modo_template = st.radio(
    "Como deseja escolher o template?",
    ["Usar template padrão", "Enviar meu template"],
    horizontal=True,
)

TEMPLATES_PADRAO = {
    "SESI": "sesi.pptx",
    "SENAI": "senai.pptx",
    "CNI": "cni.pptx",
    "IEL": "iel.pptx",
    "Sistema Indústria": "sistema_industria.pptx",
}

pasta_tamanho = "21x14.5" if formato_prisma == "21x14,5 cm" else "30x10.5"

template_file = None
template_path = None
nome_template = None
forcar_redimensionamento = False

if modo_template == "Usar template padrão":
    nome_template = st.selectbox("Escolha o modelo", list(TEMPLATES_PADRAO.keys()))
    template_path = os.path.join("templates", pasta_tamanho, TEMPLATES_PADRAO[nome_template])

    if os.path.exists(template_path):
        st.success(f"✅ Template **{nome_template}** ({formato_prisma}) carregado")
    else:
        st.error(f"❌ Arquivo `{template_path}` não encontrado. Verifique a pasta.")
else:
    template_file = st.file_uploader("Envie o template .pptx", type=["pptx"])
    if template_file:
        forcar_redimensionamento = st.checkbox(
            f"Forçar redimensionamento para {formato_prisma}", 
            value=False,
            help="⚠️ Cuidado: isso vai esticar ou comprimir o slide. Imagens e logos podem distorcer."
        )

st.markdown("---")

# BOTÃO GERAR
if st.button("🚀 GERAR PRISMAS", type="primary", use_container_width=True):
    if not nome_col and not cargo_col and not empresa_col:
        st.error("Selecione pelo menos uma coluna para mapeamento!")
        st.stop()

    try:
        progress = st.progress(0)
        status = st.empty()

        status.info("📂 Carregando template...")

        if modo_template == "Usar template padrão":
            if not template_path or not os.path.exists(template_path):
                st.error("Template padrão não encontrado.")
                st.stop()
            prs = Presentation(template_path)
        else:
            if not template_file:
                st.error("Envie o template PowerPoint!")
                st.stop()
            template_bytes = io.BytesIO(template_file.read())
            prs = Presentation(template_bytes)
            
            if forcar_redimensionamento:
                if formato_prisma == "21x14,5 cm":
                    prs.slide_width = Cm(21)
                    prs.slide_height = Cm(14.5)
                else:
                    prs.slide_width = Cm(30)
                    prs.slide_height = Cm(10.5)

        total = len(df)
        if total == 0:
            st.error("Nenhum dado no Excel!")
            st.stop()

        status.info(f"⚙️ Gerando {total} prisma(s)...")

        for i in range(total):
            duplicate_slide(prs, 0)
            progress.progress((i + 1) / (total * 2))

        campos_para_limpar = []
        if not nome_col:
            campos_para_limpar.append("NOME")
        if not cargo_col:
            campos_para_limpar.append("CARGO")
        if not empresa_col:
            campos_para_limpar.append("EMPRESA")

        for idx, (_, row) in enumerate(df.iterrows(), start=1):
            slide = prs.slides[idx]
            subs = {}
            if nome_col:
                subs["NOME"] = row[nome_col]
            if cargo_col:
                subs["CARGO"] = row[cargo_col]
            if empresa_col:
                subs["EMPRESA"] = row[empresa_col]

            substitute_text_in_slide(
                slide,
                subs,
                campos_para_limpar=campos_para_limpar,
                apagar_shape=apagar_shape_vazio,
                font_sizes=font_sizes
            )
            progress.progress((total + idx) / (total * 2))

        delete_slide(prs, 0)

        status.info("💾 Preparando arquivo PowerPoint...")
        output_pptx = io.BytesIO()
        prs.save(output_pptx)
        output_pptx.seek(0)
        pptx_bytes = output_pptx.getvalue()

        progress.progress(100)
        status.success(f"✅ {total} prisma(s) gerado(s) com sucesso!")

        st.download_button(
            label="📥 Baixar PowerPoint (.pptx)",
            data=pptx_bytes,
            file_name=f"prismas_{formato_prisma.replace(',','.').replace(' ','')}.pptx",
            mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
            use_container_width=True,
        )

    except Exception as e:
        st.error(f"❌ Erro durante a geração: {e}")
        st.exception(e)